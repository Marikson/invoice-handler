import base64
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

# XLS_FILE_PATH = "https://nordvikkft.sharepoint.com/:x:/s/Nordvik-Dokumentumok/IQB3bShMeKO-SZpYHsGexweMAYurSjFY_oj2cnTyMc74xFo?e=1dVd05"
XLS_FILE_PATH = "https://nordvikkft.sharepoint.com/:x:/s/Nordvik-Dokumentumok/IQAXzyRsieo2QYBqMGDRzZqbAVPNBm3Q9fW_P7uFTlF2-D4?e=5ANjrl"
XLS_FILE_NAME = "Szamlak_proba.xlsx"
# XLS_FILE_NAME = "Szerződés és számla_nyilvántartás_Nordvik_Kft"
GRAPH_ROOT = "https://graph.microsoft.com/v1.0"

class XlsMethods:    
    def __init__(self, session):
        self.session = session

    def _apply_cell_format(self, cell, column: int) -> None:
        if column == 5:
            cell.number_format = '# ##0 "Ft"'

        if column == 8:
            alignment = copy(cell.alignment)
            alignment.horizontal = "right"
            cell.alignment = alignment

    def _normalize_invoice_id(self, invoice_id: str | None) -> str:
        if not invoice_id:
            return ""
        return "".join(character.lower() for character in str(invoice_id) if character.isalnum())

    def _normalize_amount(self, amount) -> str:
        if amount in (None, "", "N/A"):
            return ""
        if isinstance(amount, float):
            return f"{amount:.2f}"
        return str(amount).strip()

    def _encode_share_url(self, share_url: str) -> str:
        encoded = base64.urlsafe_b64encode(share_url.encode("utf-8")).decode("ascii")
        return encoded.rstrip("=")

    def _build_share_content_url(self, share_url: str) -> str:
        encoded = self._encode_share_url(share_url)
        return f"{GRAPH_ROOT}/shares/u!{encoded}/driveItem/content"

    def _get_cached_column_values(
        self,
        file_path: str,
        sheet_name: str,
        min_row: int,
        columns: list[int],
        next_empty_row: int | None = None,
    ) -> dict[int, list[str]] | None:
        try:
            workbook_path = Path(file_path)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)

            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet '{sheet_name}' was not found in {workbook_path}")

            worksheet = workbook[sheet_name]
            cached_values = {column: [] for column in columns}
            if next_empty_row is None:
                return None

            last_data_row = next_empty_row - 1
            if last_data_row < min_row:
                return cached_values

            for row_index in range(min_row, last_data_row + 1):
                for column in columns:
                    value = worksheet.cell(row=row_index, column=column).value
                    if column == 3:
                        cached_values[column].append(self._normalize_invoice_id(value))
                    else:
                        cached_values[column].append(self._normalize_amount(value))

            return cached_values
        except Exception as e:
            print(f"Error caching Excel columns: {e}")
            return None

    def find_next_empty_row(
        self,
        file_path: str,
        sheet_name: str,
        min_row: int = 2,
        columns: list[int] | None = None,
    ) -> int | None:
        try:
            workbook_path = Path(file_path)
            workbook = load_workbook(workbook_path)

            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet '{sheet_name}' was not found in {workbook_path}")

            worksheet = workbook[sheet_name]
            target_columns = columns or [1]

            for row_index in range(worksheet.max_row, min_row - 1, -1):
                if any(worksheet.cell(row=row_index, column=column).value not in (None, "") for column in target_columns):
                    return row_index + 1

            return min_row
        except Exception as e:
            print(f"Error determining next empty Excel row: {e}")
            return None

    def invoice_id_exists(
        self,
        file_path: str,
        sheet_name: str,
        invoice_id: str | None,
        id_column: int = 3,
        min_row: int = 2,
    ) -> bool:
        try:
            normalized_invoice_id = self._normalize_invoice_id(invoice_id)
            if not normalized_invoice_id:
                return False

            workbook_path = Path(file_path)
            workbook = load_workbook(workbook_path, read_only=True)

            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet '{sheet_name}' was not found in {workbook_path}")

            worksheet = workbook[sheet_name]
            for row_index in range(min_row, worksheet.max_row + 1):
                current_value = worksheet.cell(row=row_index, column=id_column).value
                normalized_current_value = self._normalize_invoice_id(current_value)
                if not normalized_current_value:
                    continue

                if (
                    normalized_current_value == normalized_invoice_id
                    or normalized_current_value in normalized_invoice_id
                    or normalized_invoice_id in normalized_current_value
                ):
                    return True

            return False
        except Exception as e:
            print(f"Error checking invoice ID in Excel file: {e}")
            return False

    def filter_new_invoices(
        self,
        file_path: str,
        sheet_name: str,
        invoice_data: list,
        id_column: int = 3,
        amount_column: int = 5,
        min_row: int = 2,
        next_empty_row: int | None = None,
    ) -> tuple[list, list]:
        new_invoices = []
        duplicate_invoices = []

        cached_values = self._get_cached_column_values(
            file_path,
            sheet_name,
            min_row,
            [id_column, amount_column],
            next_empty_row,
        )
        if cached_values is None:
            return invoice_data, duplicate_invoices

        cached_invoice_ids = cached_values[id_column]
        cached_amounts = cached_values[amount_column]

        for invoice in invoice_data:
            normalized_invoice_id = self._normalize_invoice_id(invoice.invoice_id)
            normalized_amount = self._normalize_amount(invoice.amount)
            is_duplicate = False

            if normalized_invoice_id and normalized_amount:
                for index, cached_invoice_id in enumerate(cached_invoice_ids):
                    if cached_invoice_id != normalized_invoice_id:
                        continue

                    cached_amount = cached_amounts[index]
                    if cached_amount == normalized_amount:
                        is_duplicate = True
                        break

            if is_duplicate:
                duplicate_invoices.append(invoice)
            else:
                new_invoices.append(invoice)

        return new_invoices, duplicate_invoices

    def write_to_downloaded_xls(
        self,
        file_path: str,
        sheet_name: str,
        updates: list[dict],
    ) -> bool:
        try:
            workbook_path = Path(file_path)
            workbook = load_workbook(workbook_path)

            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet '{sheet_name}' was not found in {workbook_path}")

            worksheet = workbook[sheet_name]
            for update in updates:
                
                if "cell" in update:
                    cell = worksheet[update["cell"]]
                    cell.value = update["value"]
                    self._apply_cell_format(cell, cell.column)
                    continue

                row = update.get("row")
                column = update.get("column")
                if row is None or column is None:
                    raise ValueError("Each update must include either 'cell' or both 'row' and 'column'")

                cell = worksheet.cell(row=row, column=column, value=update.get("value"))
                self._apply_cell_format(cell, column)

            workbook.save(workbook_path)
            return True
        except Exception as e:
            print(f"Error writing to Excel file: {e}")
            return False

    def upload_xls_file(self, file_path: str, share_url: str = XLS_FILE_PATH) -> bool:
        try:
            source = Path(file_path)
            if not source.is_file():
                raise FileNotFoundError(f"Excel file not found: {source}")

            upload_url = self._build_share_content_url(share_url)
            with source.open("rb") as workbook_file:
                response = self.session.put(
                    upload_url,
                    data=workbook_file,
                    headers={
                        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    },
                    timeout=60,
                )

            response.raise_for_status()
            return True
        except Exception as e:
            print(f"Error uploading file: {e}")
            return False
    
    def download_xls_file(self, save_to: str, share_url: str = XLS_FILE_PATH) -> bool:
        try:
            url = self._build_share_content_url(share_url)
            response = self.session.get(url, timeout=60)
            response.raise_for_status()
            
            destination = Path(save_to)
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_bytes(response.content)
            
            return True
        except Exception as e:
            print(f"Error downloading file: {e}")
            return False
        
    
    def build_invoice_updates(self, invoice_data: list, start_row: int) -> list[dict]:
        updates: list[dict] = []
        for row_index, invoice in enumerate(invoice_data, start=start_row):
            updates.extend(
            [
                {"row": row_index, "column": 3, "value": invoice.invoice_id if invoice.invoice_id else "N/A"},
                {"row": row_index, "column": 4, "value": invoice.vendor if invoice.vendor else "N/A"},
                {"row": row_index, "column": 5, "value": invoice.amount if invoice.amount else "N/A"},
                {"row": row_index, "column": 8, "value": invoice.deadline.value if invoice.deadline else "N/A"},
            ]
            )
        return updates

    
    
